# -*- coding: utf-8 -*-
"""Converte `docs/defesa_roteiro.md` em `docs/defesa_roteiro.docx`.

Conversor separado do `gerar_tcc_docx.py` de propósito. Aquele produz um
documento ABNT — capa, sumário, numeração que começa na folha de rosto,
legenda acima do quadro. Este produz outra coisa: um **documento de ensaio**,
que se lê em pé, com a banca na frente.

As decisões de formatação seguem desse uso, e não de norma nenhuma:

- **uma página por slide**, para o ensaio poder ser feito folheando;
- **a fala em corpo maior e entrelinha 1,5**, destacada por barra lateral —
  é o único bloco que se lê em voz alta, e precisa ser encontrado no meio da
  página sem procurar;
- **o tempo do slide no cabeçalho**, à direita, porque é o dado que se confere
  de relance;
- **as perguntas da banca em página própria**, com resposta curta destacada da
  aprofundada: na hora, fala-se a curta.

Uso:
    python scripts/gerar_defesa_docx.py
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "docs" / "defesa_roteiro.md"
SAIDA = RAIZ / "docs" / "defesa_roteiro.docx"

FONTE = "Calibri"
FONTE_MONO = "Consolas"
CORPO_PT = 11
FALA_PT = 13          # a fala é maior: lê-se de pé, a meio metro do papel

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
CINZA_TEXTO = RGBColor(0x44, 0x44, 0x44)
VERMELHO = RGBColor(0x9B, 0x1C, 0x1C)
CINZA_FUNDO = "F2F4F7"
CINZA_CABECALHO = "E4E7EC"

_TOKENS = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")

# Rótulos que estruturam o slide. Tudo o mais que venha em negrito sozinho é
# conteúdo em destaque, não etiqueta.
_RUBRICAS = re.compile(
    r"^(Visual|Fala \(|Objetivo|Perguntas poss|Sobre |Aprofundamento|Curta)",
    re.I)
_LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")


# --------------------------------------------------------------------------
# Auxiliares de XML — python-docx não expõe sombreamento nem borda de parágrafo
# --------------------------------------------------------------------------
def _sombrear(paragrafo, cor):
    elemento = OxmlElement("w:shd")
    elemento.set(qn("w:val"), "clear")
    elemento.set(qn("w:fill"), cor)
    paragrafo._p.get_or_add_pPr().append(elemento)


def _barra_lateral(paragrafo, cor="1F3A5F", largura=18):
    """Barra à esquerda. É o que faz a fala ser achada sem ser lida."""
    pPr = paragrafo._p.get_or_add_pPr()
    bordas = OxmlElement("w:pBdr")
    esquerda = OxmlElement("w:left")
    esquerda.set(qn("w:val"), "single")
    esquerda.set(qn("w:sz"), str(largura))
    esquerda.set(qn("w:space"), "10")
    esquerda.set(qn("w:color"), cor)
    bordas.append(esquerda)
    pPr.append(bordas)


def _sombrear_celula(celula, cor):
    elemento = OxmlElement("w:shd")
    elemento.set(qn("w:val"), "clear")
    elemento.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(elemento)


def _inline(paragrafo, texto, tamanho=CORPO_PT, cor=None, italico=False):
    """Aplica negrito, itálico e monoespaçado dentro do parágrafo."""
    texto = _LINK.sub(r"\1", texto)
    for parte in _TOKENS.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = paragrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = paragrafo.add_run(parte[1:-1])
            run.font.name = FONTE_MONO
            run.font.size = Pt(tamanho - 1.5)
            if cor:
                run.font.color.rgb = cor
            continue
        elif parte.startswith("*") and parte.endswith("*"):
            run = paragrafo.add_run(parte[1:-1])
            run.italic = True
        else:
            run = paragrafo.add_run(parte)
        run.font.name = FONTE
        run.font.size = Pt(tamanho)
        if italico:
            run.italic = True
        if cor:
            run.font.color.rgb = cor


def _configurar(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(CORPO_PT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for secao in doc.sections:
        secao.top_margin = Cm(2.0)
        secao.bottom_margin = Cm(2.0)
        secao.left_margin = Cm(2.2)
        secao.right_margin = Cm(2.0)


def _cabecalho_do_slide(doc, titulo, tempo):
    """Faixa superior do slide: título à esquerda, tempo à direita."""
    tabela = doc.add_table(rows=1, cols=2)
    tabela.autofit = True
    esquerda, direita = tabela.rows[0].cells

    p = esquerda.paragraphs[0]
    run = p.add_run(titulo)
    run.bold = True
    run.font.size = Pt(15)
    run.font.name = FONTE
    run.font.color.rgb = AZUL

    p = direita.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if tempo:
        run = p.add_run(tempo)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = FONTE
        run.font.color.rgb = VERMELHO

    for celula in (esquerda, direita):
        _sombrear_celula(celula, CINZA_CABECALHO)
    doc.add_paragraph()


class Conversor:
    """Percorre o markdown uma vez, decidindo bloco a bloco."""

    def __init__(self, doc):
        self.doc = doc
        self.primeiro_slide = True

    # -- blocos -----------------------------------------------------------
    def rubrica(self, texto):
        """Rótulo de seção dentro do slide: Visual, Fala, Objetivo…"""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(texto.upper())
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.name = FONTE
        run.font.color.rgb = AZUL

    def fala(self, linhas):
        """O bloco que se lê em voz alta. Maior, espaçado e com barra."""
        texto = " ".join(l.strip() for l in linhas if l.strip())
        if not texto:
            return
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.6)
        pf.right_indent = Cm(0.3)
        pf.space_before = Pt(4)
        pf.space_after = Pt(10)
        pf.line_spacing = 1.5
        _inline(p, texto, tamanho=FALA_PT)
        _barra_lateral(p)
        _sombrear(p, CINZA_FUNDO)

    def destaque(self, texto):
        """Frase em negrito que é CONTEÚDO — a manchete visual de um slide."""
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(8)
        pf.space_after = Pt(8)
        pf.left_indent = Cm(0.4)
        _inline(p, texto, tamanho=12.5)
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = AZUL

    def paragrafo(self, texto, tamanho=CORPO_PT, cor=None, italico=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        _inline(p, texto, tamanho=tamanho, cor=cor, italico=italico)
        return p

    def item(self, texto, nivel=0):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.7 + 0.5 * nivel)
        p.paragraph_format.space_after = Pt(2)
        _inline(p, texto)

    def tabela(self, linhas):
        celulas = [[c.strip() for c in l.strip().strip("|").split("|")]
                   for l in linhas]
        celulas = [c for c in celulas
                   if not all(re.fullmatch(r":?-{2,}:?", x or "-") for x in c)]
        if not celulas:
            return
        colunas = max(len(l) for l in celulas)
        t = self.doc.add_table(rows=len(celulas), cols=colunas)
        t.style = "Table Grid"
        for i, linha in enumerate(celulas):
            for j in range(colunas):
                celula = t.rows[i].cells[j]
                celula.paragraphs[0].text = ""
                valor = linha[j] if j < len(linha) else ""
                _inline(celula.paragraphs[0], valor, tamanho=9.5)
                if i == 0:
                    _sombrear_celula(celula, CINZA_CABECALHO)
                    for run in celula.paragraphs[0].runs:
                        run.bold = True
        self.doc.add_paragraph()

    def regua(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        _barra_lateral(p, cor="D0D5DD", largura=6)


def converter():
    if not ENTRADA.is_file():
        sys.exit(f"não encontrei {ENTRADA}")

    doc = Document()
    _configurar(doc)
    conv = Conversor(doc)

    linhas = ENTRADA.read_text(encoding="utf-8").splitlines()
    i = 0
    dentro_da_fala = False
    buffer_fala = []
    contagem = {"slides": 0, "perguntas": 0, "tabelas": 0}

    while i < len(linhas):
        linha = linhas[i]
        crua = linha.rstrip()
        texto = crua.strip()

        # --- fim de um bloco de citação (a fala) -------------------------
        if dentro_da_fala and not texto.startswith(">"):
            conv.fala(buffer_fala)
            buffer_fala, dentro_da_fala = [], False

        # --- citação: acumula ------------------------------------------
        if texto.startswith(">"):
            dentro_da_fala = True
            buffer_fala.append(texto.lstrip(">").strip())
            i += 1
            continue

        # --- slide: começa página nova ----------------------------------
        m = re.match(r"^## (Slide \d+) — (.+)$", texto)
        if m:
            if not conv.primeiro_slide:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            conv.primeiro_slide = False
            contagem["slides"] += 1
            # o tempo está na rubrica **Fala (X)** logo adiante
            tempo = ""
            for j in range(i, min(i + 40, len(linhas))):
                t = re.search(r"\*\*Fala \(([^)]+)\)\*\*", linhas[j])
                if t:
                    tempo = t.group(1)
                    break
            _cabecalho_do_slide(doc, f"{m.group(1)} · {m.group(2)}", tempo)
            i += 1
            continue

        # --- demais títulos ---------------------------------------------
        m = re.match(r"^(#{1,3}) (.+)$", texto)
        if m:
            nivel, titulo = len(m.group(1)), m.group(2)
            if nivel == 1:
                if not conv.primeiro_slide:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                conv.primeiro_slide = False
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(12)
                run = p.add_run(titulo)
                run.bold = True
                run.font.size = Pt(19)
                run.font.name = FONTE
                run.font.color.rgb = AZUL
            elif nivel == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                run = p.add_run(titulo)
                run.bold = True
                run.font.size = Pt(14)
                run.font.name = FONTE
                run.font.color.rgb = AZUL
            else:
                if re.match(r"^\d+\.", titulo):
                    contagem["perguntas"] += 1
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(titulo)
                run.bold = True
                run.font.size = Pt(12)
                run.font.name = FONTE
                run.font.color.rgb = AZUL
            i += 1
            continue

        # --- linha inteiramente em negrito -------------------------------
        # Nem toda linha em negrito é rótulo: o slide 10 abre com uma FRASE em
        # negrito, e tratá-la como rubrica a encolhia para maiúsculas de nove
        # pontos — conteúdo virando etiqueta. Só é rubrica o que está na lista.
        m = re.match(r"^\*\*([^*]+)\*\*$", texto)
        if m:
            rotulo = m.group(1)
            if _RUBRICAS.match(rotulo):
                conv.rubrica(rotulo)
            else:
                conv.destaque(rotulo)
            i += 1
            continue

        # --- tabela ------------------------------------------------------
        if texto.startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i])
                i += 1
            conv.tabela(bloco)
            contagem["tabelas"] += 1
            continue

        # --- régua -------------------------------------------------------
        if re.fullmatch(r"-{3,}", texto):
            conv.regua()
            i += 1
            continue

        # --- lista -------------------------------------------------------
        m = re.match(r"^(\s*)[-*] (.+)$", crua)
        if m:
            conv.item(m.group(2), nivel=len(m.group(1)) // 2)
            i += 1
            continue

        # --- parágrafo ---------------------------------------------------
        if texto:
            italico = texto.startswith("*") and texto.endswith("*") \
                and "**" not in texto
            conv.paragrafo(texto,
                           cor=CINZA_TEXTO if italico else None,
                           italico=False)
        i += 1

    if dentro_da_fala:
        conv.fala(buffer_fala)

    doc.save(SAIDA)
    print(f"gerado:  {SAIDA}")
    print(f"tamanho: {SAIDA.stat().st_size:,} bytes")
    print(f"blocos:  {contagem['slides']} slides, "
          f"{contagem['perguntas']} perguntas, {contagem['tabelas']} tabelas")
    print("Uma página por slide; a fala está no bloco com barra azul à esquerda.")


if __name__ == "__main__":
    converter()
